"""Ingest Blufab-format files so at the event you just drop their data and run.

Expected inputs (put them in a folder, e.g. data_in/):
  - panel_features.csv     : one row per panel, the features read from the drawing/BOM
  - micro_operations.xlsx  : matrix — micro-operation per row, panel per column, time in cells
                             (this is exactly how Blufab's example-times Excel is laid out)
  - production_orders.csv   : order_id, panels (semicolon-separated)

The ONLY thing to adjust when the real files arrive is the label mapping below
(NAME_MAP) — e.g. the Portuguese micro-operation names — and the CSV column names.
"""
import csv

from .features import counts, OP_LABEL

# Map the micro-operation labels AS THEY APPEAR in the Excel -> our internal keys.
# Auto-built from our English labels; add the real (Portuguese) names at the event.
NAME_MAP = {v: k for k, v in OP_LABEL.items()}
NAME_MAP.update({
    # --- Blufab's real (Portuguese) labels map here when the Excel arrives ---
    "Montar gabarito": "setup_jig",
    "Colocar canal": "place_track",          # canais = tracks
    "Colocar montante": "place_stud",        # montantes = studs
    "Colocar travessa": "place_noggin",
    "Cravar / aparafusar estrutura": "clinch_frame",
    "Colocar placa": "place_board",
    "Aparafusar placa": "screw_board",
    "Furar": "drill_hole",
    "Colocar cerâmica": "place_ceramic",
    "Virar, etiquetar e paletizar": "turn_and_finish",
    "Desaparafusar (erro)": "unscrew_exception",
})


def load_features_csv(path):
    out = []
    with open(path, newline="", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            out.append({
                "name": r["name"],
                "n_frames": int(r["n_frames"]),
                "n_studs": int(r["n_studs"]),
                "n_noggins": int(r["n_noggins"]),
                "n_boards": int(r["n_boards"]),
                "height_mm": int(r["height_mm"]),
                "n_drillings": int(r.get("n_drillings", 0) or 0),
                "cladding": (r.get("cladding") or "standard"),
            })
    return out


def load_microops_xlsx(path):
    """Read the micro-op-by-panel time matrix. Returns {panel_name: {op_key: seconds}}."""
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True).active
    header = [c.value for c in ws[1]]
    panel_names = header[1:]
    out = {p: {} for p in panel_names}
    for row in ws.iter_rows(min_row=2, values_only=True):
        key = NAME_MAP.get(row[0])
        if not key:
            continue  # unknown label -> flag at the event and add to NAME_MAP
        for i, p in enumerate(panel_names):
            v = row[i + 1]
            if v not in (None, "", 0):
                out[p][key] = float(v)
    return out


def load_orders_csv(path):
    orders = []
    with open(path, newline="", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            orders.append({"order_id": r["order_id"],
                           "panels": [x.strip() for x in r["panels"].split(";") if x.strip()]})
    return orders


def build_panels(features, microops):
    """Combine drawing features + measured times into the model's panel format.
    A panel is 'train' if it has measured times in the Excel, else 'test'."""
    panels = []
    for spec in features:
        p = dict(spec)
        p["counts"] = counts(spec)
        ot = microops.get(spec["name"], {})
        if ot:
            p["op_times_s"] = ot
            p["total_time_s"] = round(sum(ot.values()), 1)
            p["split"] = "train"
        else:
            p["split"] = "test"
        panels.append(p)
    return panels
